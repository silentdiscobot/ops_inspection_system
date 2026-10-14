# -*- coding: utf-8 -*-
"""服务器与服务信息的自选字段 Excel 导出。"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from excel_export_style import (
    FONT_NAME, TEXT_COLOR, add_empty_state, add_styled_table,
    configure_print_layout, fit_wrapped_row, prepare_export_sheet,
    safe_excel_value, style_status_cell, style_table_header,
)


# 字段列表是导出的唯一白名单。凭据、密钥、口令及其密文字段不得加入。
SERVER_EXPORT_FIELDS = (
    {'key': 'name', 'label': '名称', 'width': 20, 'default': True},
    {'key': 'ip', 'label': 'IP地址', 'width': 17, 'default': True},
    {'key': 'physical_ip', 'label': '物理机IP', 'width': 17, 'default': True},
    {'key': 'port', 'label': '端口', 'width': 10, 'default': True},
    {'key': 'username', 'label': '登录用户名', 'width': 16, 'default': True},
    {'key': 'auth_type', 'label': '认证方式', 'width': 14, 'default': True},
    {'key': 'resource_type', 'label': '资源类型', 'width': 14, 'default': True},
    {'key': 'group_names', 'label': '所属环境', 'width': 24, 'default': True},
    {'key': 'os_type', 'label': '系统类型', 'width': 16, 'default': True},
    {'key': 'rack_number', 'label': '机柜编号', 'width': 14, 'default': False},
    {'key': 'room_affiliation', 'label': '机房归属', 'width': 18, 'default': False},
    {'key': 'rack_level', 'label': '机柜层级', 'width': 14, 'default': False},
    {'key': 'service_count', 'label': '关联服务数', 'width': 13, 'default': False},
    {'key': 'notes', 'label': '备注', 'width': 36, 'default': True},
)

SERVICE_EXPORT_FIELDS = (
    {'key': 'name', 'label': '服务名称', 'width': 22, 'default': True},
    {'key': 'ip', 'label': 'IP地址', 'width': 17, 'default': True},
    {'key': 'server_name', 'label': '服务器名称', 'width': 20, 'default': True},
    {'key': 'os_type', 'label': '系统类型', 'width': 16, 'default': False},
    {'key': 'system_names', 'label': '归属系统', 'width': 24, 'default': True},
    {'key': 'environment_names', 'label': '归属环境', 'width': 24, 'default': True},
    {'key': 'service_type', 'label': '服务类型', 'width': 14, 'default': True},
    {'key': 'run_mode', 'label': '运行方式', 'width': 14, 'default': True},
    {'key': 'version', 'label': '版本号', 'width': 14, 'default': True},
    {'key': 'ports', 'label': '端口信息', 'width': 24, 'default': True},
    {'key': 'status', 'label': '服务状态', 'width': 13, 'default': True},
    {'key': 'owner', 'label': '负责人', 'width': 14, 'default': True},
    {'key': 'auto_start', 'label': '开机自启', 'width': 12, 'default': False},
    {'key': 'account_count', 'label': '账号数量', 'width': 12, 'default': False},
    {'key': 'deploy_dir', 'label': '部署目录', 'width': 34, 'default': False},
    {'key': 'config_path', 'label': '配置文件路径', 'width': 36, 'default': False},
    {'key': 'log_path', 'label': '日志路径', 'width': 34, 'default': False},
    {'key': 'access_url', 'label': '访问地址', 'width': 36, 'default': False},
    {'key': 'start_command', 'label': '启动命令', 'width': 48, 'default': False},
    {'key': 'stop_command', 'label': '停止命令', 'width': 48, 'default': False},
    {'key': 'restart_command', 'label': '重启命令', 'width': 48, 'default': False},
    {'key': 'notes', 'label': '备注', 'width': 36, 'default': False},
    {'key': 'created_by', 'label': '录入人', 'width': 14, 'default': False},
    {'key': 'created_at', 'label': '创建时间', 'width': 20, 'default': False},
    {'key': 'updated_at', 'label': '更新时间', 'width': 20, 'default': False},
)


def select_export_fields(field_specs, requested):
    """按白名单和用户顺序返回字段；未知字段必须拒绝。"""
    field_map = {item['key']: item for item in field_specs}
    requested = [str(item or '').strip() for item in (requested or []) if str(item or '').strip()]
    if not requested:
        requested = [item['key'] for item in field_specs if item.get('default')]
    selected = []
    seen = set()
    for key in requested:
        if key not in field_map:
            raise ValueError('包含不允许导出的字段')
        if key not in seen:
            selected.append(field_map[key])
            seen.add(key)
    if not selected:
        raise ValueError('请至少选择一个导出字段')
    return selected


def _cell_value(field, value):
    key = field['key']
    if key == 'auth_type':
        value = 'SSH私钥' if value == 'key' else '密码'
    elif key == 'auto_start':
        value = '是' if bool(value) else '否'
    elif value == '无':
        value = ''
    if key in ('created_at', 'updated_at') and value:
        try:
            return datetime.strptime(str(value), '%Y-%m-%d %H:%M:%S')
        except ValueError:
            pass
    return safe_excel_value(value if value is not None else '')


def build_asset_export_excel(title, sheet_name, rows, selected_fields, generated_at=None):
    """生成适合屏幕查看和打印的单表资产清单。"""
    generated_at = generated_at or datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    column_count = len(selected_fields)
    last_column = prepare_export_sheet(
        ws, title, column_count, generated_at, len(rows), '当前筛选结果',
    )

    ws.append([])
    ws.append([item['label'] for item in selected_fields])
    header_row = 4
    style_table_header(ws, header_row, column_count)

    for item in rows:
        ws.append([_cell_value(field, item.get(field['key'])) for field in selected_fields])

    if rows:
        last_row = ws.max_row
    else:
        last_row = header_row + 1
        add_empty_state(ws, last_row, last_column)
    add_styled_table(
        ws, 'AssetExportTable', header_row, last_column, last_row, bool(rows),
    )

    for index, field in enumerate(selected_fields, start=1):
        ws.column_dimensions[get_column_letter(index)].width = field.get('width', 18)

    wrap_keys = {
        'notes', 'deploy_dir', 'config_path', 'log_path', 'access_url',
        'start_command', 'stop_command', 'restart_command', 'ports',
        'group_names', 'system_names', 'environment_names',
    }
    center_keys = {
        'port', 'auth_type', 'resource_type', 'os_type', 'rack_number',
        'service_count', 'service_type', 'run_mode', 'version', 'status',
        'owner', 'auto_start', 'account_count', 'created_by',
        'created_at', 'updated_at',
    }
    grid_side = Side(style='thin', color='D3DEE6')
    light_border = Border(
        left=grid_side, right=grid_side, top=grid_side, bottom=grid_side,
    )
    data_rows = ws.iter_rows(min_row=5, max_row=last_row) if rows else ()
    for row_number, row in enumerate(data_rows, start=5):
        wrapped = []
        for index, cell in enumerate(row):
            field = selected_fields[index]
            key = field['key']
            is_wrapped = key in wrap_keys
            cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
            cell.alignment = Alignment(
                horizontal='center' if key in center_keys else 'left',
                vertical='top' if is_wrapped else 'center',
                wrap_text=is_wrapped,
            )
            cell.border = light_border
            if key in ('created_at', 'updated_at') and isinstance(cell.value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            elif key in ('port', 'service_count', 'account_count') and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
            if key in ('status', 'auto_start'):
                style_status_cell(cell)
            if key == 'access_url' and str(cell.value or '').lower().startswith(('http://', 'https://')):
                cell.hyperlink = str(cell.value)
                cell.font = Font(name=FONT_NAME, size=10, color='176FA8', underline='single')
            if is_wrapped:
                wrapped.append((cell.value, field.get('width', 18)))
        fit_wrapped_row(ws, row_number, wrapped)

    configure_print_layout(
        ws, title, header_row, last_column, last_row, column_count,
        freeze_first_column=column_count > 8,
    )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
