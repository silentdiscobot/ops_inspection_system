# -*- coding: utf-8 -*-
"""运维清单类 Excel 的统一阅读样式。"""
import math

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo


FONT_NAME = '微软雅黑'
TITLE_COLOR = '173B57'
TEXT_COLOR = '263B49'
MUTED_COLOR = '6B7F8F'
HEADER_COLOR = '1F4E78'
HEADER_BORDER = '6F96B5'
TAB_COLOR = '16834B'

STATUS_STYLES = {
    '运行中': ('E7F6EC', '16733A'),
    '已停止': ('FFF4D6', '946200'),
    '已下线': ('EEF1F4', '52616D'),
    '已匹配': ('E7F6EC', '16733A'),
    '已确认正常': ('E7F6EC', '16733A'),
    '已恢复关闭': ('E7F6EC', '16733A'),
    '待执行': ('E9F2FB', '1F5F91'),
    '已过期': ('FFF0E5', 'A95417'),
    '已取消': ('EEF1F4', '52616D'),
    '异常': ('FDEAEA', 'B42318'),
    '已确认异常': ('FDEAEA', 'B42318'),
    '恢复待确认': ('FFF4D6', '946200'),
    '是': ('E7F6EC', '16733A'),
    '否': ('F1F4F6', '52616D'),
}


def safe_excel_value(value):
    """阻止用户输入在 Excel 中被解释为公式。"""
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def prepare_export_sheet(ws, title, column_count, generated_at, record_count,
                         scope_text='当前筛选结果'):
    """创建紧凑、易读的标题和元数据区域。"""
    from openpyxl.utils import get_column_letter

    last_column = get_column_letter(column_count)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = TAB_COLOR
    if column_count > 1:
        ws.merge_cells(f'A1:{last_column}1')
        ws.merge_cells(f'A2:{last_column}2')
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT_NAME, size=16, bold=True, color=TITLE_COLOR)
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 30
    ws['A2'] = (
        f'数据范围：{scope_text}    记录数：{record_count:,}    '
        f'生成时间：{generated_at.strftime("%Y-%m-%d %H:%M:%S")}'
    )
    ws['A2'].font = Font(name=FONT_NAME, size=9, color=MUTED_COLOR, italic=True)
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 21
    ws.row_dimensions[3].height = 9
    return last_column


def style_table_header(ws, header_row, column_count):
    white_separator = Side(style='thin', color=HEADER_BORDER)
    bottom = Side(style='medium', color='163B5C')
    header_cells = next(ws.iter_rows(
        min_row=header_row, max_row=header_row,
        min_col=1, max_col=column_count,
    ))
    for cell in header_cells:
        cell.fill = PatternFill('solid', fgColor=HEADER_COLOR)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=white_separator, right=white_separator, bottom=bottom)
    ws.row_dimensions[header_row].height = 28


def add_styled_table(ws, table_name, header_row, last_column, last_row, has_rows):
    if has_rows:
        table = Table(displayName=table_name, ref=f'A{header_row}:{last_column}{last_row}')
        table.tableStyleInfo = TableStyleInfo(
            name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f'A{header_row}:{last_column}{header_row}'


def add_empty_state(ws, row_number, last_column):
    if last_column != 'A':
        ws.merge_cells(f'A{row_number}:{last_column}{row_number}')
    cell = ws.cell(row_number, 1, '当前筛选条件下暂无数据')
    cell.font = Font(name=FONT_NAME, size=10, color=MUTED_COLOR, italic=True)
    cell.fill = PatternFill('solid', fgColor='F7F9FB')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    outline = Side(style='thin', color='D3DEE6')
    cell.border = Border(left=outline, right=outline, top=outline, bottom=outline)
    ws.row_dimensions[row_number].height = 38


def style_status_cell(cell):
    style = STATUS_STYLES.get(str(cell.value or '').strip())
    if not style:
        return
    fill, color = style
    cell.fill = PatternFill('solid', fgColor=fill)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=color)
    cell.alignment = Alignment(horizontal='center', vertical='center')


def fit_wrapped_row(ws, row_number, cells_with_width, minimum=23, maximum=72):
    """按长文本粗略估算行高，避免整列被截断或所有行过高。"""
    lines = 1
    for value, width in cells_with_width:
        text = str(value or '')
        if not text:
            continue
        effective_width = max(8, int(width * 0.9))
        estimated = sum(max(1, math.ceil(len(part) / effective_width)) for part in text.splitlines())
        lines = max(lines, estimated)
    ws.row_dimensions[row_number].height = min(maximum, max(minimum, 16 * lines + 7))


def configure_print_layout(ws, title, header_row, last_column, last_row, column_count,
                           freeze_first_column=False):
    ws.freeze_panes = f'B{header_row + 1}' if freeze_first_column else f'A{header_row + 1}'
    ws.print_title_rows = f'{header_row}:{header_row}'
    ws.print_area = f'A1:{last_column}{last_row}'
    ws.page_setup.orientation = 'landscape' if column_count > 6 else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 3 if column_count > 20 else (2 if column_count > 12 else 1)
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.55, bottom=0.55, header=0.2, footer=0.25,
    )
    ws.oddHeader.left.text = title
    ws.oddHeader.left.size = 9
    ws.oddHeader.left.color = MUTED_COLOR
    ws.oddFooter.center.text = '第 &[Page] 页 / 共 &[Pages] 页'
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.color = MUTED_COLOR
