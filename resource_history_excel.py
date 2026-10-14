# -*- coding: utf-8 -*-
"""资源配置变更历史 Excel 导出。"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from excel_export_style import (
    FONT_NAME, TEXT_COLOR, add_empty_state, add_styled_table,
    configure_print_layout, fit_wrapped_row, prepare_export_sheet,
    safe_excel_value, style_status_cell, style_table_header,
)


STATUS_LABELS = {
    'matched': '已匹配', 'expired': '已过期', 'cancelled': '已取消',
    'confirmed_normal': '已确认正常', 'resolved': '已恢复关闭',
    'planned': '待执行', 'confirmed_abnormal': '已确认异常',
    'recovered_pending': '恢复待确认',
}


def _as_datetime(value):
    try:
        return datetime.strptime(str(value or ''), '%Y-%m-%d %H:%M:%S')
    except (TypeError, ValueError):
        return None


def build_resource_history_excel(rows, generated_at=None):
    wb = Workbook()
    ws = wb.active
    ws.title = '变更历史'
    generated_at = generated_at or datetime.now()
    last_column = prepare_export_sheet(
        ws, '资源配置变更历史', 10, generated_at, len(rows), '当前筛选结果',
    )

    headers = ['时间', '服务器名称', '服务器IP', '记录类型', '变化内容',
               '状态', '理由/备注', '操作人', '变更单号', '记录编号']
    ws.append([])
    ws.append(headers)
    header_row = 4
    style_table_header(ws, header_row, 10)

    for item in rows:
        ws.append([
            _as_datetime(item.get('occurred_at')) or str(item.get('occurred_at') or ''),
            safe_excel_value(item.get('server_name') or ''),
            safe_excel_value(item.get('server_ip') or ''),
            safe_excel_value(item.get('record_type') or ''),
            safe_excel_value(item.get('summary') or ''),
            STATUS_LABELS.get(item.get('status'), item.get('status') or ''),
            safe_excel_value(item.get('reason') or ''),
            safe_excel_value(item.get('operator') or ''),
            safe_excel_value(item.get('ticket_no') or ''),
            safe_excel_value(item.get('record_id') or ''),
        ])

    if rows:
        last_row = ws.max_row
    else:
        last_row = header_row + 1
        add_empty_state(ws, last_row, last_column)
    add_styled_table(
        ws, 'ResourceChangeHistory', header_row, last_column, last_row, bool(rows),
    )
    widths = {'A': 20, 'B': 20, 'C': 17, 'D': 13, 'E': 54,
              'F': 15, 'G': 42, 'H': 14, 'I': 20, 'J': 12}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    grid_side = Side(style='thin', color='D3DEE6')
    light_border = Border(
        left=grid_side, right=grid_side, top=grid_side, bottom=grid_side,
    )
    data_rows = ws.iter_rows(min_row=5, max_row=last_row) if rows else ()
    for row_number, row in enumerate(data_rows, start=5):
        for cell in row:
            is_wrapped = cell.column in (5, 7)
            cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
            cell.alignment = Alignment(
                horizontal='center' if cell.column in (1, 3, 4, 6, 8, 9, 10) else 'left',
                vertical='top' if is_wrapped else 'center', wrap_text=is_wrapped,
            )
            cell.border = light_border
        row[0].number_format = 'yyyy-mm-dd hh:mm:ss'
        if row[3].value == '计划变更':
            row[3].fill = PatternFill('solid', fgColor='E9F2FB')
            row[3].font = Font(name=FONT_NAME, size=10, bold=True, color='1F5F91')
        elif row[3].value == '实际变化':
            row[3].fill = PatternFill('solid', fgColor='FFF0E5')
            row[3].font = Font(name=FONT_NAME, size=10, bold=True, color='A95417')
        style_status_cell(row[5])
        fit_wrapped_row(ws, row_number, [(row[4].value, 54), (row[6].value, 42)])

    configure_print_layout(ws, '资源配置变更历史', header_row, last_column, last_row, 10)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
